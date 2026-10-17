"""Read-only workbook parser (Phase 1).

Loads the workbook twice with openpyxl — ``data_only=False`` (formulas) and
``data_only=True`` (cached values) — and builds ``CellRecord``s. Never
executes formulas or macros.

Golden-rule notes:
- Workbooks written by openpyxl (including the test fixtures) have **no**
  cached values. Everything must work without them: when no formula cell has
  a cached value, the parser sets the workbook-level warning
  ``cached_values_missing`` and leaves ``value`` as None.
- ``.xlsm``/macro workbooks are rejected with a clear error in v1.
- ``MAX_FILE_MB`` and ``MAX_CELLS`` guards from ``app.config`` are enforced.

``cells`` holds only non-empty records keyed by ``"Sheet!A1"``; an absent key
means an empty cell.
"""

from __future__ import annotations

import os
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula

from ...config import MAX_CELLS, MAX_FILE_MB
from ..models import CellRecord, CellRef, ParsedWorkbook
from .labels import attach_labels
from .refs import cell_key

CACHED_VALUES_MISSING = "cached_values_missing"

_MACRO_SUFFIXES = {".xlsm", ".xltm", ".xlsb", ".xlam", ".xla"}


def _classify(value: Any, formula: str | None) -> str | None:
    """Return the CellRecord kind, or None for empty (not stored)."""
    if formula is not None:
        return "formula"
    if value is None:
        return None
    if isinstance(value, str):
        return None if not value.strip() else "label"
    if isinstance(value, bool):
        return "input"
    if isinstance(value, (int, float, datetime)):
        return "input"
    # Anything else with content (rare types) is treated as an input value.
    return "input"


def _defined_names(workbook: Any) -> dict[str, str]:
    """Map defined name -> target text, skipping internal ``_xlnm.*`` names."""
    names: dict[str, str] = {}
    try:
        items = list(workbook.defined_names.values())
    except Exception:
        return names
    for dn in items:
        name = getattr(dn, "name", None)
        target = getattr(dn, "attr_text", None)
        if not name or not target:
            continue
        if name.lower().startswith("_xlnm."):
            continue
        names.setdefault(name, target)
    return names


def read_workbook(path: str | os.PathLike[str]) -> ParsedWorkbook:
    """Parse ``path`` into a ``ParsedWorkbook``. Read-only; never executes."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in _MACRO_SUFFIXES or suffix == ".xlsm":
        raise ValueError(
            f"Macro workbooks ({path.suffix}) are unsupported in v1; "
            "please save a macro-free .xlsx copy."
        )
    try:
        size_mb = path.stat().st_size / (1024 * 1024)
    except OSError as exc:
        raise ValueError(f"Cannot read workbook: {path}") from exc
    if size_mb > MAX_FILE_MB:
        raise ValueError(
            f"Workbook is {size_mb:.1f} MB, above the {MAX_FILE_MB} MB limit."
        )

    try:
        wb_formulas = load_workbook(path, data_only=False)
    except Exception as exc:
        raise ValueError(
            f"Cannot parse workbook {path.name!r}: not a readable .xlsx file "
            f"({exc})."
        ) from exc
    try:
        wb_values = load_workbook(path, data_only=True)
    except Exception:
        wb_values = None

    # Workbook-size guard on the used bounding boxes.
    total_box = 0
    for ws in wb_formulas.worksheets:
        total_box += max(ws.max_row, 1) * max(ws.max_column, 1)
        if total_box > MAX_CELLS:
            raise ValueError(
                f"Workbook exceeds the {MAX_CELLS}-cell limit."
            )

    defined = _defined_names(wb_formulas)
    cells: dict[str, CellRecord] = {}
    formulas_seen = 0
    cached_seen = 0

    value_sheets = {}
    if wb_values is not None:
        try:
            value_sheets = {ws.title: ws for ws in wb_values.worksheets}
        except Exception:
            value_sheets = {}

    for ws in wb_formulas.worksheets:
        title = ws.title
        ws_values = value_sheets.get(title)
        for row in ws.iter_rows():
            for cell in row:
                raw = cell.value
                if isinstance(raw, ArrayFormula):
                    # Array formulas read back as objects; their text is the
                    # formula (with "=" prefix per openpyxl's writer).
                    text = raw.text or ""
                    raw = text if text.startswith("=") else f"={text}"
                formula = raw if isinstance(raw, str) and raw.startswith("=") else None
                cached: Any = None
                if ws_values is not None:
                    try:
                        cached = ws_values[cell.coordinate].value
                        if isinstance(cached, str) and cached.startswith("="):
                            cached = None  # data_only mirror missing: no cache
                    except Exception:
                        cached = None
                value = cached if formula is not None else raw
                kind = _classify(raw, formula)
                if kind is None:
                    continue
                if formula is not None:
                    formulas_seen += 1
                    if cached is not None:
                        cached_seen += 1
                number_format = (
                    cell.number_format if cell.number_format else None
                )
                key = cell_key(title, cell.coordinate)
                cells[key] = CellRecord(
                    ref=CellRef(sheet=title, address=cell.coordinate),
                    formula=formula,
                    value=value,
                    number_format=number_format,
                    kind=kind,  # type: ignore[arg-type]
                )

    attach_labels(cells)

    warnings: list[str] = []
    has_cached = True
    if formulas_seen and cached_seen == 0:
        warnings.append(CACHED_VALUES_MISSING)
        has_cached = False

    return ParsedWorkbook(
        cells=cells,
        defined_names=defined,
        warnings=warnings,
        has_cached_values=has_cached,
        source_path=str(path),
    )
