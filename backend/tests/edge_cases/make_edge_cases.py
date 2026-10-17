"""Edge-case workbook generator for QA Stage 3.

Builds a set of small, focused ``.xlsx`` files exercising parser/detector
corners. Real Excel-cached values and error cells (which openpyxl cannot
write) are injected by filling the ``<v>`` nodes openpyxl itself emits.
Call :func:`build_all` with a directory; it returns ``{name: Path}``.
"""

from __future__ import annotations

import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _inject_cached(path: Path, values: dict[str, tuple[str, str | None]]) -> Path:
    """Fill cached values into an openpyxl-written file.

    ``values`` maps ``"Sheet!A1"`` -> ``(text, error_or_None)``. Returns a
    new ``*_cached.xlsx`` path next to ``path`` (original untouched).
    """
    with zipfile.ZipFile(path) as zin:
        names = zin.namelist()
        sheets = {
            name: zin.read(name)
            for name in names
            if name.startswith("xl/worksheets/sheet")
        }
        others = {name: zin.read(name) for name in names if name not in sheets}
    by_sheet: dict[str, dict[str, tuple[str, str | None]]] = {}
    for key, spec in values.items():
        sheet, _, address = key.partition("!")
        by_sheet.setdefault(sheet, {})[address] = spec
    titles = list(_sheet_titles(path))
    ET.register_namespace("", NS)
    for name, data in sheets.items():
        idx = int(name.rsplit("sheet", 1)[1].split(".")[0]) - 1
        title = titles[idx] if idx < len(titles) else None
        mapping = by_sheet.get(title or "", {})
        if not mapping:
            continue
        root = ET.fromstring(data)
        for cell in root.iter(f"{{{NS}}}c"):
            spec = mapping.get(cell.get("r", ""))
            if spec is None:
                continue
            text, error = spec
            if error is not None:
                cell.set("t", "e")
            node = cell.find(f"{{{NS}}}v")
            if node is None:
                node = ET.SubElement(cell, f"{{{NS}}}v")
            node.text = error if error is not None else text
        sheets[name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    out = path.with_name(path.stem + "_cached.xlsx")
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in {**others, **sheets}.items():
            zout.writestr(name, data)
    return out


def _sheet_titles(path: Path) -> list[str]:
    from openpyxl import load_workbook

    return load_workbook(path, read_only=True).sheetnames


def build_all(out_dir: str | Path) -> dict[str, Path]:
    """Write every edge-case workbook; return ``{name: path}``."""
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    made: dict[str, Path] = {}

    def save(wb: Workbook, name: str) -> Path:
        path = out / f"{name}.xlsx"
        wb.save(path)
        made[name] = path
        return path

    # 1. Empty workbook (single untouched sheet).
    save(Workbook(), "empty")

    # 2. One empty sheet plus a label-only sheet.
    wb = Workbook()
    wb.active.title = "Only"
    notes = wb.create_sheet("Notes")
    notes["A1"] = "hello"
    save(wb, "blank_sheet")

    # 3. Awkward sheet names + quoted cross-sheet refs.
    wb = Workbook()
    wb.active.title = "My Sheet"
    wb.active["A1"] = 1
    obrien = wb.create_sheet("O'Brien")
    obrien["A1"] = "='My Sheet'!A1*2"
    uni = wb.create_sheet("Ünïcödé 日本語")
    uni["A1"] = "='My Sheet'!A1+1"
    party = wb.create_sheet("🎉 Party")
    party["A1"] = "=SUM('My Sheet'!A1:A1)"
    save(wb, "weird_names")

    # 4. Defined names (workbook-level; sheet-scoped attempted too).
    wb = Workbook()
    wb.active.title = "Data"
    wb.active["A1"] = 3
    wb.active["B1"] = "=Rate*2"
    wb.defined_names["Rate"] = DefinedName("Rate", attr_text="Data!$A$1")
    try:
        wb.defined_names["Local"] = DefinedName(
            "Local", attr_text="Data!$A$1", localSheetId=0
        )
    except Exception:
        pass
    save(wb, "names")

    # 5. Whole-column / whole-row refs (must collapse, never expand).
    wb = Workbook()
    wb.active.title = "S"
    wb.active["A1"] = 1
    wb.active["A2"] = 2
    wb.active["A3"] = 3
    wb.active["B1"] = "=SUM(A:A)"
    wb.active["B2"] = "=SUM(1:3)"
    save(wb, "whole_col")

    # 6. Tricky literals: strings with commas/parens, numbers in strings,
    #    excluded function args, plus one genuine violation (control).
    wb = Workbook()
    wb.active.title = "S"
    wb.active["A1"] = 2
    wb.active["B1"] = '="a,b(c)"&A1'
    wb.active["B2"] = '="Item 100"'
    wb.active["B3"] = "=ROUND(A1*B1,2)"
    wb.active["B4"] = "=MATCH(9,A1:A2,0)"
    wb.active["B5"] = "=VLOOKUP(A1,A1:A2,2,TRUE)"
    wb.active["B6"] = "=INDEX(A1:A2,1,1)"
    wb.active["B7"] = "=IF(A1>0,1,0)"
    wb.active["B8"] = "=A1*1.15"
    save(wb, "tricky")

    # 7. Merged cells, hidden sheet, hidden row.
    wb = Workbook()
    wb.active.title = "S"
    wb.active.merge_cells("A1:B1")
    wb.active["A1"] = "Total"
    wb.active["C1"] = "=A2*2"
    wb.active["A2"] = 21
    wb.active.row_dimensions[2].hidden = True
    meta = wb.create_sheet("Meta")
    meta.sheet_state = "hidden"
    meta["A1"] = 7
    save(wb, "hidden_merged")

    # 8. External workbook links (must be ignored, never crash).
    wb = Workbook()
    wb.active.title = "S"
    wb.active["A1"] = "=[1]Sheet1!A1"
    wb.active["A2"] = "='C:\\dir\\[b.xlsx]S'!A1"
    save(wb, "external")

    # 9. Errors: injected cached error + #REF! literal + typed error string.
    wb = Workbook()
    wb.active.title = "S"
    wb.active["A1"] = 5
    wb.active["A2"] = "=A1/0"
    wb.active["A3"] = "=SUM(#REF!)"
    wb.active["A4"] = "#N/A"
    path = save(wb, "errors")
    made["errors_cached"] = _inject_cached(
        path, {"S!A2": ("", "#DIV/0!")}
    )

    # 10. Cached values present (no errors).
    wb = Workbook()
    wb.active.title = "S"
    wb.active["A1"] = 5
    wb.active["A2"] = "=A1*2"
    path = save(wb, "cached")
    made["cached_ok"] = _inject_cached(path, {"S!A2": ("10", None)})

    # 11. Cycles: size 2, size 5, self-loop.
    wb = Workbook()
    wb.active.title = "S"
    wb.active["A1"] = "=A2+1"
    wb.active["A2"] = "=A1*2"
    wb.active["B1"] = "=B2+1"
    wb.active["B2"] = "=B3+1"
    wb.active["B3"] = "=B4+1"
    wb.active["B4"] = "=B5+1"
    wb.active["B5"] = "=B1+1"
    wb.active["C1"] = "=C1+1"
    save(wb, "loops")

    # 12. 5,000-row filled column with one sandwiched constant.
    wb = Workbook()
    wb.active.title = "S"
    wb.active["A1"] = "V"
    wb.active["B1"] = "W"
    for r in range(2, 5002):
        wb.active[f"B{r}"] = r
        wb.active[f"A{r}"] = f"=B{r}*2"
    wb.active["A2502"] = 7.7
    save(wb, "sandwich_5k")

    # 13. Percent scale: 8 vs 0.08 in one percent-formatted column.
    wb = Workbook()
    wb.active.title = "S"
    wb.active["A1"] = "Rate"
    wb.active["A2"] = 8
    wb.active["A3"] = 0.08
    wb.active["A2"].number_format = "0%"
    wb.active["A3"].number_format = "0%"
    save(wb, "percent")

    # 14. Array formula (reads back as an object, not a string).
    wb = Workbook()
    wb.active.title = "Data"
    wb.active["A1"] = 1
    wb.active["A2"] = 2
    wb.active["B1"] = ArrayFormula("B1", "=SUM(A1:A2*2)")
    save(wb, "array")

    return made


if __name__ == "__main__":
    import sys

    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("files")
    for name, path in build_all(target).items():
        print(f"{name}: {path}")
