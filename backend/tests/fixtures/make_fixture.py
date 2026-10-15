"""Fixture generator (Phase 1): ``dirty.xlsx``, ``clean.xlsx``.

Builds both workbooks exactly per spec section 11.1 with openpyxl, plus
``expected_issues.json`` listing the expected ``{detector, sheet, address}``
for ``dirty.xlsx`` — six entries (the five planted problems plus the extra
orphan ``Settings!B4``, which becomes unreferenced in ``dirty.xlsx`` when
``Summary!B5`` stops referencing it).

``dirty.xlsx`` differs from ``clean.xlsx`` in exactly five places:
1. ``Summary!B5`` = ``=B4*1.15`` (hardcoded_constant)
2. ``Pricing!C13`` = the number ``99.5`` (pattern_break)
3. ``Summary!B3`` = ``=SUM(Pricing!C2:C20)`` (range_gap)
4. ``Settings!A3/B3`` = "Unused discount" / ``0.10`` (orphan_input)
5. ``Summary!B7/B8`` loops (circular)

Workbooks written by openpyxl carry no cached values, and the whole tool must
work without them. Recalculation via LibreOffice is therefore opt-in only:
set ``SHEETSENSE_RECALC=1`` and have ``soffice`` on PATH. Tests pass without it.

Usage: ``python make_fixture.py [out_dir]`` (defaults to this directory).
"""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook

N_ROWS = 20
FIRST_ROW = 2
LAST_ROW = FIRST_ROW + N_ROWS - 1  # 21

EXPECTED_ISSUES = [
    {"detector": "hardcoded_constant", "sheet": "Summary", "address": "B5"},
    {"detector": "pattern_break", "sheet": "Pricing", "address": "C13"},
    {"detector": "range_gap", "sheet": "Summary", "address": "B3"},
    {"detector": "orphan_input", "sheet": "Settings", "address": "B3"},
    {"detector": "orphan_input", "sheet": "Settings", "address": "B4"},
    {"detector": "circular", "sheet": "Summary", "address": "B7"},
]


def _build_base() -> Workbook:
    """Shared layout for clean and dirty (dirty edits applied afterwards)."""
    wb = Workbook()

    inputs = wb.active
    inputs.title = "Inputs"
    inputs["A1"] = "Item"
    inputs["B1"] = "Unit Cost"
    inputs["C1"] = "Qty"
    for i in range(1, N_ROWS + 1):
        r = FIRST_ROW + i - 1
        inputs[f"A{r}"] = f"Item-{i:02d}"
        inputs[f"B{r}"] = 5.0 * i
        inputs[f"C{r}"] = (i % 5) + 1

    settings = wb.create_sheet("Settings")
    settings["A1"] = "Markup"
    settings["B1"] = 0.25
    settings["A2"] = "Tax rate"
    settings["B2"] = 0.08
    settings["A4"] = "Surcharge"
    settings["B4"] = 0.15

    pricing = wb.create_sheet("Pricing")
    pricing["A1"] = "Item"
    pricing["B1"] = "Cost Total"
    pricing["C1"] = "Price"
    pricing["D1"] = "With Tax"
    for r in range(FIRST_ROW, LAST_ROW + 1):
        pricing[f"A{r}"] = f"=Inputs!A{r}"
        pricing[f"B{r}"] = f"=Inputs!B{r}*Inputs!C{r}"
        pricing[f"C{r}"] = f"=B{r}*(1+Settings!$B$1)"
        pricing[f"D{r}"] = f"=C{r}+C{r}*Settings!$B$2"

    summary = wb.create_sheet("Summary")
    summary["A2"] = "Total cost"
    summary["B2"] = "=SUM(Pricing!B2:B21)"
    summary["A3"] = "Total price"
    summary["B3"] = "=SUM(Pricing!C2:C21)"
    summary["A4"] = "Total with tax"
    summary["B4"] = "=SUM(Pricing!D2:D21)"
    summary["A5"] = "Total with surcharge"
    summary["B5"] = "=B4*(1+Settings!B4)"

    return wb


def _apply_dirty(wb: Workbook) -> None:
    settings = wb["Settings"]
    pricing = wb["Pricing"]
    summary = wb["Summary"]

    # 1. hardcoded constant instead of the surcharge reference.
    summary["B5"] = "=B4*1.15"
    # 2. constant inside a formula column.
    pricing["C13"] = 99.5
    # 3. range skipping the last data row.
    summary["B3"] = "=SUM(Pricing!C2:C20)"
    # 4. unused labeled input (also strands Settings!B4 — see module docstring).
    settings["A3"] = "Unused discount"
    settings["B3"] = 0.10
    # 5. two-cell reference loop.
    summary["A7"] = "Loop A"
    summary["B7"] = "=B8+1"
    summary["A8"] = "Loop B"
    summary["B8"] = "=B7*2"


def _maybe_recalc(path: Path) -> None:
    """Optionally recalculate cached values via LibreOffice (opt-in)."""
    if os.environ.get("SHEETSENSE_RECALC") != "1":
        return
    soffice = shutil.which("soffice")
    if soffice is None:
        print("SHEETSENSE_RECALC=1 but soffice not found; skipping recalc.")
        return
    subprocess.run(
        [soffice, "--headless", "--convert-to", "xlsx", "--outdir",
         str(path.parent), str(path)],
        check=True,
    )


def make_fixtures(out_dir: str | Path | None = None) -> Path:
    """Write clean.xlsx, dirty.xlsx and expected_issues.json; return out_dir."""
    out = Path(out_dir) if out_dir is not None else Path(__file__).resolve().parent
    out.mkdir(parents=True, exist_ok=True)

    clean_path = out / "clean.xlsx"
    _build_base().save(clean_path)
    _maybe_recalc(clean_path)

    dirty_path = out / "dirty.xlsx"
    wb = _build_base()
    _apply_dirty(wb)
    wb.save(dirty_path)
    _maybe_recalc(dirty_path)

    with open(out / "expected_issues.json", "w", encoding="utf-8") as fh:
        json.dump(EXPECTED_ISSUES, fh, indent=2)
        fh.write("\n")
    return out


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else None
    print(f"Fixtures written to {make_fixtures(target)}")
